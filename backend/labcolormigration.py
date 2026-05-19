import os
import sys
import firebase_admin
from firebase_admin import credentials, firestore
from openpyxl import load_workbook

SERVICE_ACCOUNT_KEY = "serviceAccountKey.json"
EXCEL_PATH          = "Chromascope_file.xlsx"

PRODUCT_SHEETS = {
    "Lip":        "Products — Lip",
    "Face":       "Products — Face",
    "Cheek":      "Products — Cheek",
    "Eye":        "Products — Eye",
    "Concealer":  "Products — Concealer",
    "Multi-Use":  "Products — Multi-Use",
}

def clean_str(val) -> str:
    if val is None: return ""
    return str(val).strip()

def init_firebase():
    if not os.path.exists(SERVICE_ACCOUNT_KEY):
        print(f"[ERROR] {SERVICE_ACCOUNT_KEY} not found.")
        sys.exit(1)
    cred = credentials.Certificate(SERVICE_ACCOUNT_KEY)
    firebase_admin.initialize_app(cred)
    return firestore.client()

def migrate_color_data(wb, db_client):
    print("\n🚀 Starting Color Data Migration...")
    total = 0

    for category, sheet_name in PRODUCT_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            print(f"  [SKIP] Sheet '{sheet_name}' not found.")
            continue

        ws = wb[sheet_name]
        batch = db_client.batch()
        count = 0

        # Data starts on row 5 based on your Excel format
        for row in ws.iter_rows(min_row=5):
            pid = row[0].value
            if not pid or not str(pid).startswith("PH-"):
                continue

            pid_clean = clean_str(pid)
            doc_ref = db_client.collection("products").document(pid_clean)

            # Mapping columns to the specific Color Engine indices found in your file
            data = {
                "product_id":          pid_clean,
                "brand":               clean_str(row[1].value),
                "product_name":        clean_str(row[2].value),
                "category":            category,
                "season_tags":         clean_str(row[13].value),
                "hex_color":           clean_str(row[14].value) or "#FFFFFF",
                "lab_L":               float(row[15].value or 0),
                "lab_a":               float(row[16].value or 0),
                "lab_b":               float(row[17].value or 0),
                "season_tag_computed": clean_str(row[19].value),
            }

            # .set with merge=True ensures we don't delete existing image_urls
            batch.set(doc_ref, data, merge=True)
            count += 1
            total += 1

            # Commit in chunks to respect Firestore limits
            if count % 400 == 0:
                batch.commit()
                batch = db_client.batch()

        batch.commit()
        print(f"  ✓ {category}: {count} products updated.")

    print(f"\n✅ Migration Complete. Total products processed: {total}")

if __name__ == "__main__":
    if not os.path.exists(EXCEL_PATH):
        print(f"[ERROR] Excel file '{EXCEL_PATH}' not found in root.")
        sys.exit(1)

    db_client = init_firebase()
    # Read-only and data-only to ensure we get the computed values of formulas
    wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    
    migrate_color_data(wb, db_client)