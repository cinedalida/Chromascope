import os
import sys
import firebase_admin
from firebase_admin import credentials, firestore, storage
from openpyxl import load_workbook

SERVICE_ACCOUNT_KEY = "serviceAccountKey.json"
FIREBASE_PROJECT_ID = "chromascopedatabase"
STORAGE_BUCKET      = "chromascopedatabase.firebasestorage.app"
EXCEL_PATH          = "Chromascope_file.xlsx"
IMAGE_FOLDER        = "images"

PRODUCT_SHEETS = {
    "Lip":        "Products — Lip",
    "Face":       "Products — Face",
    "Cheek":      "Products — Cheek",
    "Eye":        "Products — Eye",
    "Concealer":  "Products — Concealer",
    "Multi-Use":  "Products — Multi-Use",
}

PROFILE_SHEETS = {
    "Lip":        "Profiles — Lip",
    "Face":       "Profiles — Face",
    "Cheek":      "Profiles — Cheek",
    "Eye":        "Profiles — Eye",
    "Concealer":  "Profiles — Concealer",
    "Multi-Use":  "Profiles — Multi-Use",
}

INGREDIENT_SHEET = "Ingredient Safety DB"
SYNERGY_SHEET    = "Synergy Pairs"

def clean_str(val) -> str:
    if val is None:
        return ""
    return str(val).strip()

def bool_val(val) -> bool:
    if isinstance(val, bool):
        return val
    if isinstance(val, str):
        return val.strip().upper() in ("TRUE", "YES", "Y", "1")
    return False

def safe_id(text: str) -> str:
    return (
        text.strip()
            .lower()
            .replace(" ", "_")
            .replace("/", "_")
            .replace("(", "")
            .replace(")", "")
            .replace(",", "")
            .replace(".", "")
            [:500]
    )

def commit_batch(batch, db_client):
    batch.commit()
    return db_client.batch()

def init_firebase():
    if not os.path.exists(SERVICE_ACCOUNT_KEY):
        print(f"\n[ERROR] serviceAccountKey.json not found.")
        sys.exit(1)

    cred = credentials.Certificate(SERVICE_ACCOUNT_KEY)
    firebase_admin.initialize_app(cred, {"storageBucket": STORAGE_BUCKET})
    print(f"[Firebase] Connected to project: {FIREBASE_PROJECT_ID}")
    return firestore.client(), storage.bucket()

def migrate_ingredients(wb, db_client):
    print("\n[1/5] Migrating Ingredient Safety DB...")

    if INGREDIENT_SHEET not in wb.sheetnames:
        print(f"  [SKIP] Sheet '{INGREDIENT_SHEET}' not found.")
        return

    ws    = wb[INGREDIENT_SHEET]
    batch = db_client.batch()
    count = 0

    for row in ws.iter_rows(min_row=4):
        inci = row[0].value
        if not inci or not clean_str(inci):
            continue

        doc_id  = safe_id(clean_str(inci))
        doc_ref = db_client.collection("ingredients").document(doc_id)

        data = {
            "inci_name":               clean_str(inci),
            "common_aliases":          clean_str(row[1].value),
            "comedogenic_rating":      float(row[2].value or 0),
            "safety_flag_oily":        bool_val(row[3].value),
            "safety_flag_dry":         bool_val(row[4].value),
            "safety_flag_combination": bool_val(row[5].value),
            "safety_flag_acne_prone":  bool_val(row[6].value),
            "safety_flag_normal":      bool_val(row[7].value),
            "severity":                clean_str(row[8].value).lower() or "low",
            "concern_category":        clean_str(row[9].value),
            "notes":                   clean_str(row[10].value),
            "source":                  clean_str(row[11].value),
        }

        batch.set(doc_ref, data)
        count += 1

        if count % 490 == 0:
            batch = commit_batch(batch, db_client)
            print(f"  Batch committed ({count} so far)...")

    batch.commit()
    print(f"  ✓ Ingredients migrated: {count}")

def migrate_products(wb, db_client):
    print("\n[2/5] Migrating Products...")
    total = 0

    for category, sheet_name in PRODUCT_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            print(f"  [SKIP] Sheet '{sheet_name}' not found.")
            continue

        ws    = wb[sheet_name]
        batch = db_client.batch()
        count = 0

        for row in ws.iter_rows(min_row=5):
            pid = row[0].value
            if not pid or not str(pid).strip().startswith("PH-"):
                continue

            pid_clean   = clean_str(pid)
            cat_folder  = category.lower().replace("-", "_").replace(" ", "_")
            doc_ref     = db_client.collection("products").document(pid_clean)

            data = {
                "product_id":          pid_clean,
                "brand":               clean_str(row[1].value),
                "product_name":        clean_str(row[2].value),
                "sub_category":        clean_str(row[3].value),
                "category":            category,
                "where_to_buy":        clean_str(row[4].value),
                "price":               clean_str(row[5].value),
                "full_inci_list":      clean_str(row[6].value),
                "key_actives":         clean_str(row[7].value),
                "has_fragrance":       clean_str(row[8].value),
                "has_alcohol":         clean_str(row[9].value),
                "paraben_free":        clean_str(row[10].value),
                "comedogenic_risk":    clean_str(row[11].value),
                "flagged_summary":     clean_str(row[12].value),
                "season_tags":         clean_str(row[13].value),
                "hex_color":           clean_str(row[14].value),
                "season_tag_computed": clean_str(row[19].value),
                "image_storage_path":  f"products/{cat_folder}/{pid_clean}.jpg",
                "image_url":           "",
            }

            batch.set(doc_ref, data)
            count += 1
            total += 1

            if count % 490 == 0:
                batch = commit_batch(batch, db_client)

        batch.commit()
        print(f"  ✓ {category}: {count} products")

    print(f"  ✓ Total products migrated: {total}")

def migrate_profiles(wb, db_client):
    print("\n[3/5] Migrating Safety Profiles...")
    total = 0

    for category, sheet_name in PROFILE_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            continue

        ws    = wb[sheet_name]
        batch = db_client.batch()
        count = 0

        for row in ws.iter_rows(min_row=5):
            pid = row[0].value
            if not pid or not str(pid).strip().startswith("PH-"):
                continue

            pid_clean = clean_str(pid)
            doc_ref   = db_client.collection("safety_profiles").document(pid_clean)

            data = {
                "product_id":            pid_clean,
                "product_name":          clean_str(row[1].value),
                "safe_for_oily":         bool_val(row[2].value),
                "safe_for_dry":          bool_val(row[3].value),
                "safe_for_combination":  bool_val(row[4].value),
                "safe_for_acne_prone":   bool_val(row[5].value),
                "safe_for_normal":       bool_val(row[6].value),
                "flagged_ingredients":   clean_str(row[7].value),
                "highest_severity":      clean_str(row[8].value).lower() or "low",
                "filter_out_oily":       bool_val(row[9].value),
                "filter_out_acne_prone": bool_val(row[10].value),
                "notes_for_dev":         clean_str(row[11].value),
            }

            batch.set(doc_ref, data)
            count += 1
            total += 1

            if count % 490 == 0:
                batch = commit_batch(batch, db_client)

        batch.commit()
        print(f"  ✓ {category}: {count} profiles")

    print(f"  ✓ Total profiles migrated: {total}")

def migrate_synergy_pairs(wb, db_client):
    print("\n[4/5] Migrating Synergy Pairs...")

    if SYNERGY_SHEET not in wb.sheetnames:
        default_pairs = [
            {
                "pair_id":                "SYN-001",
                "ingredient_a":           "Parfum",
                "ingredient_b":           "Alcohol Denat.",
                "ingredient_a_inci":      "parfum",
                "ingredient_b_inci":      "alcohol denat",
                "verdict":                "excluded",
                "severity":               "high",
                "reason":                 "Fragrance + Alcohol — dual irritant.",
                "applies_to_skin_types":  "sensitive, dry, rosacea, eczema",
                "source":                 "Komericki et al. 2025; Paula's Choice",
            },
            {
                "pair_id":                "SYN-002",
                "ingredient_a":           "Cocos Nucifera Oil",
                "ingredient_b":           "Isopropyl Myristate",
                "ingredient_a_inci":      "cocos nucifera",
                "ingredient_b_inci":      "isopropyl myristate",
                "verdict":                "excluded",
                "severity":               "high",
                "reason":                 "Combined comedogenic load very high.",
                "applies_to_skin_types":  "oily, acne_prone, combination",
                "source":                 "INCIDecoder; CosDNA",
            }
        ]

        batch = db_client.batch()
        for pair in default_pairs:
            doc_ref = db_client.collection("synergy_pairs").document(pair["pair_id"])
            batch.set(doc_ref, pair)
        batch.commit()
        return

    ws    = wb[SYNERGY_SHEET]
    batch = db_client.batch()
    count = 0

    for row in ws.iter_rows(min_row=2):
        pair_id = row[0].value
        if not pair_id:
            continue

        doc_ref = db_client.collection("synergy_pairs").document(clean_str(pair_id))
        data = {
            "pair_id":               clean_str(row[0].value),
            "ingredient_a":          clean_str(row[1].value),
            "ingredient_b":          clean_str(row[2].value),
            "ingredient_a_inci":     clean_str(row[3].value),
            "ingredient_b_inci":     clean_str(row[4].value),
            "verdict":               clean_str(row[5].value),
            "severity":              clean_str(row[6].value),
            "reason":                clean_str(row[7].value),
            "applies_to_skin_types": clean_str(row[8].value),
            "source":                clean_str(row[9].value),
        }

        batch.set(doc_ref, data)
        count += 1

    batch.commit()
    print(f"  ✓ Synergy pairs migrated: {count}")

def upload_images(db_client, bkt):
    print(f"\n[5/5] Uploading images from '{IMAGE_FOLDER}/'...")

    if not os.path.exists(IMAGE_FOLDER):
        print(f"  [SKIP] Folder '{IMAGE_FOLDER}' not found.")
        return

    uploaded = 0
    skipped  = 0

    for category in ["lip", "face", "cheek", "eye", "concealer", "multi_use"]:
        folder = os.path.join(IMAGE_FOLDER, category)
        if not os.path.exists(folder):
            continue

        for filename in os.listdir(folder):
            ext = filename.rsplit(".", 1)[-1].lower()
            if ext not in ("jpg", "jpeg", "png", "webp"):
                continue

            pid          = filename.rsplit(".", 1)[0]
            local_path   = os.path.join(folder, filename)
            storage_path = f"products/{category}/{filename}"

            blob = bkt.blob(storage_path)
            content_types = {"jpg":"image/jpeg","jpeg":"image/jpeg","png":"image/png","webp":"image/webp"}
            blob.upload_from_filename(local_path, content_type=content_types.get(ext, "image/jpeg"))
            blob.make_public()
            image_url = blob.public_url

            try:
                db_client.collection("products").document(pid).update({
                    "image_url":          image_url,
                    "image_storage_path": storage_path,
                })
                uploaded += 1
            except Exception as e:
                skipped += 1

    print(f"  ✓ Images uploaded: {uploaded} | Skipped: {skipped}")

def verify_migration(db_client):
    print("\n[Verify] Checking Firestore counts...")
    collections = ["ingredients", "products", "safety_profiles", "synergy_pairs"]
    for col in collections:
        docs = list(db_client.collection(col).limit(1000).stream())
        print(f"  {col}: {len(docs)} documents")

if __name__ == "__main__":
    if not os.path.exists(EXCEL_PATH):
        sys.exit(1)

    db_client, bkt = init_firebase()

    wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    
    migrate_ingredients(wb, db_client)
    migrate_products(wb, db_client)
    migrate_profiles(wb, db_client)
    migrate_synergy_pairs(wb, db_client)
    upload_images(db_client, bkt)
    verify_migration(db_client)

    print("\n✓ Migration complete!\n")